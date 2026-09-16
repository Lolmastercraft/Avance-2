"""Genera el documento de evidencias a partir de resultados reales guardados."""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.7)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(8)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(27)
doc.styles['Heading 1'].font.size = Pt(20)
doc.styles['Heading 2'].font.size = Pt(14)
doc.styles['Caption'].font.size = Pt(9)
doc.core_properties.title = 'Evidencias del marketplace Mercado Nube'
doc.core_properties.author = 'Luis Alfonso Juárez Amaro'


def p(text, style=None):
    return doc.add_paragraph(text, style)


def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)


def picture(name, caption, max_height=5.7):
    target = ROOT / name
    with Image.open(target) as im:
        width = min(7.1, max_height * im.width / im.height)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(3)
    run = para.add_run()
    inline = run.add_picture(str(target), width=Inches(width))
    inline._inline.docPr.set('descr', caption)
    p(caption, 'Caption')


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + key), value)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate([headers, *rows]):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j, (cell, text) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[j])
            cell.text = str(text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                item = OxmlElement('w:' + edge)
                item.set(qn('w:w'), '90')
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            props.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '24483C' if i == 0 else ('F1F4F2' if i % 2 else 'FFFFFF'))
            props.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    if i == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        if i == 0:
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        t.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    p('')


red = json.loads((ROOT / 'reportes/roja/veredicto.json').read_text(encoding='utf-8'))
green = json.loads((ROOT / 'reportes/verde/veredicto.json').read_text(encoding='utf-8'))
sbom = json.loads((ROOT / 'reportes/sbom_cyclonedx.json').read_text(encoding='utf-8'))

doc.add_heading('Evidencias del marketplace Mercado Nube', 0)
p('Avance del proyecto 2', 'Subtitle')
p('Luis Alfonso Juárez Amaro\nVerificación del 16 de septiembre de 2026')
p('Mercado Nube permite registrar usuarios, publicar productos con imágenes, agregar artículos al carrito y confirmar pedidos. La API y el servicio de notificaciones funcionan en contenedores separados sobre EC2; los datos se guardan en RDS y las imágenes en S3. Este documento reúne las evidencias funcionales, de infraestructura y del pipeline de seguridad.')
p('Repositorio privado: https://github.com/Lolmastercraft/Avance-2\nAplicación: https://100.24.2.141\nEl evaluador necesita permiso de lectura del repositorio. La aplicación depende de que el laboratorio siga activo y utiliza un certificado HTTPS autofirmado.')
picture('docs/diagrama_arquitectura.png', 'Figura 1. Arquitectura implementada, servicios propios, controles de red y endpoint /salud.', 4.25)
p('Resultado técnico: aplicación probada contra AWS real y evidencias roja y verde del mismo pipeline. El video de demostración queda a cargo del alumno; su guion se incluye en docs/guion_video.md.')

page('Catálogo y carrito')
p('La sesión de prueba registró un usuario, consultó el catálogo y añadió un producto. El servidor calcula importes en centavos y comprueba existencias; no confía en precios enviados por el navegador.')
picture('docs/capturas/01_catalogo.png', 'Figura 2. Catálogo real con imágenes leídas desde el bucket privado mediante la API.', 4.6)
picture('docs/capturas/02_carrito.png', 'Figura 3. Carrito de la sesión autenticada antes de confirmar la compra.', 2.65)

page('Pedido confirmado y publicación')
p('La compra crea pedido, renglones y mensaje de salida en una transacción. Un segundo contenedor procesa el mensaje y registra una confirmación simulada. No se cobra dinero ni se envían correos reales.')
picture('docs/capturas/03_pedido_confirmado.png', 'Figura 4. Pedido persistido y confirmación procesada por el servicio de notificaciones.', 3.3)
p('La publicación posterior utilizó una imagen PNG enviada por formulario. El servidor la validó, recodificó como JPEG y guardó en S3 bajo una clave única; los metadatos del producto se almacenaron en RDS.')
picture('docs/capturas/05_producto_s3.png', 'Figura 5. Producto creado desde la aplicación después de guardar la imagen en S3.', 3.3)

page('Contenedores y conexión a la base de datos')
p('Docker Compose mantiene dos servicios propios, API y notificaciones, y un proxy Nginx. Los servicios propios usan UID 10001, sistema de archivos de solo lectura, capacidades eliminadas y un chequeo de salud. Solo el proxy publica los puertos del host.')
picture('docs/capturas/07_contenedores.png', 'Figura 6. Visualización de la salida real de Systems Manager y consulta SQL a RDS.', 5.4)
p('La consulta de pg_stat_ssl confirmó TLS 1.3 y el usuario marketplace_runtime, distinto del administrador. Los conteos de registros prueban el uso real de la base de datos. La salida íntegra está en reportes/aws/containers_rds.txt.')
doc.add_heading('Comprobación de salud', 2)
p('El endpoint /salud verifica SELECT 1 en RDS y HeadBucket en S3. La etapa live exige respuesta correcta y valida el certificado TLS con la copia pública del certificado del laboratorio, sin desactivar esa comprobación.')

page('Recursos de S3 y RDS en AWS')
p('RDS está cifrada y no es pública. Su grupo de seguridad acepta PostgreSQL únicamente desde el grupo de seguridad de EC2. Los dos buckets tienen bloqueo público completo, cifrado AES256, versionado y una política que rechaza transporte sin TLS.')
picture('docs/capturas/06_aws_s3_rds.png', 'Figura 7. Datos reales consultados por AWS API. No es una captura de la consola de AWS.', 6.35)
p('Fuente: reportes/aws/infraestructura.json. El bucket de productos contiene las seis ilustraciones iniciales y la imagen publicada durante la prueba. El bucket de evidencias almacena los paquetes de despliegue. En el video se mostrarán los recursos directamente en la consola de AWS.')

page('Bloqueo y aprobación del pipeline')
p('Ambas corridas ejecutan ocho etapas con los mismos umbrales. El candidato rojo añade, en una copia temporal aislada, una función que concatena entrada en un comando shell. Bandit detecta el riesgo de inyección B602 de severidad alta y bloquea. Ese código no se ejecuta ni llega a EC2.')
picture('docs/capturas/08_pipeline.png', 'Figura 8. Resultados registrados por las herramientas y decisión integrada de cada corrida.', 5.4)
p('La versión corregida excluye la función insegura y mantiene búsquedas parametrizadas mediante SQLAlchemy. Todos los controles deben aprobar; una herramienta ausente, un error o una etapa omitida impide PERMITIR. scripts/deploy.py exige el veredicto verde y la huella de los archivos desplegables.')
p(f"SBOM: CycloneDX {sbom['specVersion']} con {len(sbom['components'])} componentes del entorno de producción. Los registros completos, reportes JSON, pruebas JUnit y SBOM están en reportes/. Umbrales y exclusiones: docs/tabla_decisiones_pipeline.md.")

page('Autoevaluación de requisitos mínimos')
table(['Requisito', 'Resultado y evidencia'], [
('Marketplace Python', 'Cumple: catálogo, carrito, pedidos y publicación; capturas 01 a 05.'),
('Registro e inicio de sesión', 'Cumple: contraseñas scrypt, sesiones, CSRF y pruebas automatizadas.'),
('Dos contenedores propios', 'Cumple: API y notificaciones; Compose y evidencia de ejecución.'),
('Pieza distintiva', 'Cumple: confirmación simulada de pedidos, outbox e idempotencia.'),
('S3 real privado y cifrado', 'Cumple: lectura y subida real; configuración consultada por API.'),
('RDS real privada y cifrada', 'Cumple: datos persistidos, TLS y acceso solo desde EC2.'),
('/salud', 'Cumple: usado por HEALTHCHECK, pipeline y diagrama.'),
('Infraestructura como código', 'Cumple: infra/*.tf y módulo Terraform; Checkov y políticas propias.'),
('Endurecimiento y secretos', 'Cumple: UID no root, base fija, HEALTHCHECK; secretos fuera de Git.'),
('Pipeline integrado', f"Cumple: roja {red['verdict']} y verde {green['verdict']}; mismos umbrales."),
('SBOM y documentación', 'Cumple: CycloneDX, README, ADR, tabla de decisiones y declaración de IA.'),
('Video de 3 a 5 minutos', 'Pendiente de grabación y enlace por el alumno; guion incluido.'),
], [2.25,4.85])
p('Esta autoevaluación describe evidencia técnica y no asigna una calificación. La evaluación final corresponde al docente. El repositorio es privado: se debe habilitar acceso al evaluador antes de enviar el enlace.')

page('Reproducción y límites del laboratorio')
doc.add_heading('Cómo reproducir la evidencia', 2)
p('El README explica cómo preparar por separado las dependencias de producción y las herramientas, ejecutar las pruebas, generar las dos corridas y validar la entrega. pipeline/demostrar.py crea la copia insegura temporal, conserva el reporte rojo y después evalúa las fuentes corregidas. El nombre de la corrida no determina su resultado.')
p('Las fuentes están en app/, la infraestructura en infra/terraform/, el pipeline en pipeline/ y los scripts operativos en scripts/. Los reportes incluyen fecha UTC y huellas de fuentes. El verificador propio verificar_entrega.sh comprueba archivos y veredictos; con --sin-video excluye únicamente el pendiente audiovisual.')
doc.add_heading('Decisiones y límites aceptados', 2)
p('El laboratorio usa una EC2 y RDS Single AZ para limitar consumo; no ofrece alta disponibilidad. HTTPS emplea un certificado autofirmado y la dirección pública puede cambiar al reiniciar EC2. El notificador simula la entrega y no integra pagos ni proveedores de correo.')
p('La cuenta Academy restringe ciertas operaciones IAM y S3. Se usa el rol autorizado por el laboratorio, con más permisos que los deseables en producción. Por una denegación institucional de la API Object Lock, Terraform crea los buckets mediante una operación idempotente y administra sus controles con recursos declarativos. La excepción y sus efectos de limpieza están documentados en el ADR.')
p('El análisis de dependencias cubre los paquetes Python de producción, no las vulnerabilidades del sistema operativo de las imágenes. Tampoco se certifican DAST exhaustivo, cumplimiento de licencias, firma de artefactos o resistencia a carga. La tabla del pipeline explica estas exclusiones; una corrida verde no equivale a ausencia total de riesgos.')
doc.add_heading('Uso de inteligencia artificial', 2)
p('Codex participó sustancialmente en la implementación, configuración, pruebas, diagnóstico y documentación. El alumno aportó el tema Marketplace, materiales, cuenta de laboratorio y repositorio. No se atribuyen al alumno correcciones manuales que realizó la herramienta. Antes de presentar, debe revisar el código y explicar las decisiones; la declaración completa está en docs/declaracion_uso_ia.md.')
doc.add_heading('Cierre de la entrega', 2)
p('Grabar el video en el orden exigido: acción completa, contenedores, recursos S3 y RDS en la consola y pipeline rojo y verde. Agregar su URL a docs/enlace_video.txt, conceder acceso al evaluador y subir este documento a la plataforma. No mostrar contraseñas, parámetros descifrados, archivos de estado o variables de entorno durante la grabación.')

out = ROOT / 'entrega/Evidencias_Avance2_Mercado_Nube.docx'
out.parent.mkdir(exist_ok=True)
doc.save(out)
print(out)
